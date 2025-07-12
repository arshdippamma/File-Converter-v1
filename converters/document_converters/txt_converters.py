import os
from fpdf import FPDF
from docx import Document

def txt_to_pdf(input_path):
    output_path = os.path.splitext(input_path)[0] + ".pdf"

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    with open(input_path, 'r', encoding='utf-8') as f:
        for line in f:
            pdf.cell(200, 10, txt=line.strip(), ln=True)

    pdf.output(output_path)

    return output_path

def txt_to_docx(input_path):
    output_path = os.path.splitext(input_path)[0] + ".docx"

    doc = Document()
    with open(input_path, 'r', encoding='utf-8') as f:
        for line in f:
            doc.add_paragraph(line.strip())
    
    doc.save(output_path)

    return output_path